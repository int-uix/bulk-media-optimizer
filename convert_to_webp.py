#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Minimal UI to bulk-convert a directory tree to .webp,
optionally resizing images to a max width while preserving aspect ratio.
Also copies video files (e.g., .mp4) into the destination, preserving structure.

Changes vs your current script:
- NO "remember last folders" behavior (removed config save/load completely).
- Destination is auto-set immediately after selecting a source:
    "<source_name> webp" next to the source folder.
  Example: /.../boho  ->  /.../boho webp
  If it already exists: "boho webp_2", "boho webp_3", ...
- Destination can still be changed manually or via Browse.
- XLSX index includes BOTH:
    * converted images (non-webp -> webp)
    * copied .webp images (webp -> webp copy)
  (Videos are copied but not indexed.)

Requirements:
  python3 -m pip install pillow openpyxl
Optional for drag & drop:
  python3 -m pip install tkinterdnd2

Run:
  python3 webp_resizer_ui.py
"""

import os
import sys
import time
import shutil
import subprocess
import threading
from pathlib import Path
from typing import Iterable, Optional, List, Tuple

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# Optional drag & drop support
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    TkBase = TkinterDnD.Tk
    DND_AVAILABLE = True
except ImportError:
    DND_FILES = None  # type: ignore[assignment]
    TkBase = tk.Tk  # fallback, no drag & drop
    DND_AVAILABLE = False

try:
    from PIL import Image, ImageOps
except ImportError:
    raise SystemExit(
        "Pillow is not installed.\nInstall it with:\n  python3 -m pip install pillow"
    )

try:
    from openpyxl import Workbook
except ImportError:
    raise SystemExit(
        "openpyxl is not installed.\nInstall it with:\n  python3 -m pip install openpyxl"
    )


# Image formats to process; .webp is included so we can COPY it without recompressing
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".gif", ".webp"}
# Video formats to copy over as-is
VIDEO_EXTS = {".mp4", ".mov", ".avi", ".mkv"}


# -----------------------
# File helpers
# -----------------------
def iter_files(root: Path) -> Iterable[Path]:
    for dirpath, _, filenames in os.walk(root):
        for name in filenames:
            yield Path(dirpath) / name


def is_image(path: Path) -> bool:
    return path.suffix.lower() in IMAGE_EXTS


def is_video(path: Path) -> bool:
    return path.suffix.lower() in VIDEO_EXTS


def is_webp(path: Path) -> bool:
    return path.suffix.lower() == ".webp"


# -----------------------
# Image helpers
# -----------------------
def resize_preserving_aspect(img: Image.Image, max_width: int) -> Image.Image:
    img = ImageOps.exif_transpose(img)  # respect EXIF orientation
    w, h = img.size
    if max_width <= 0 or w <= max_width:
        return img
    new_w = max_width
    new_h = int(round(h * (new_w / w)))
    return img.resize((new_w, new_h), resample=Image.Resampling.LANCZOS)


def convert_to_webp(
    src: Path,
    dst: Path,
    max_width: Optional[int],
    quality: int = 85,
    lossless: bool = False,
) -> None:
    dst.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(src) as im:
        if max_width is None:
            im = ImageOps.exif_transpose(im)
        else:
            im = resize_preserving_aspect(im, max_width)

        # Normalize modes for better WEBP results
        if im.mode == "P":
            im = im.convert("RGBA")
        elif im.mode in ("CMYK", "YCbCr"):
            im = im.convert("RGB")

        save_kwargs = {
            "format": "WEBP",
            "quality": quality,
            # auto-lossless if alpha present (preserve transparency)
            "lossless": lossless or ("A" in im.getbands()),
            "method": 6,
            "icc_profile": im.info.get("icc_profile"),
        }
        if "exif" in im.info:
            save_kwargs["exif"] = im.info["exif"]

        im.save(dst, **save_kwargs)


def copy_file(src: Path, dst: Path) -> None:
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)  # preserves timestamps & metadata where possible


def write_xlsx(rows: List[Tuple[str, str]], out_path: Path) -> None:
    """
    rows: list of (filename_no_ext, ColumnB_value)
    out_path: path to .xlsx file to save
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Images"

    # headers
    ws["A1"] = "filename"
    ws["B1"] = "path"

    # data
    for i, (name, p) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=p)

    # widths
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 100

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def make_auto_destination_for_source(src: Path) -> Path:
    """
    Create a destination path next to source folder:
      "<src.name> webp", "<src.name> webp_2", "<src.name> webp_3", ...
    (Does NOT create it on disk here.)
    """
    base_dir = src.parent
    base_name = f"{src.name} webp"
    candidate = base_dir / base_name
    idx = 2
    while candidate.exists():
        candidate = base_dir / f"{base_name}_{idx}"
        idx += 1
    return candidate


# -----------------------
# UI App
# -----------------------
class App(TkBase):
    def __init__(self):
        super().__init__()
        self.title("WEBP Converter (Minimal)")
        self.geometry("700x520")
        self.resizable(False, False)

        # main state vars
        self.src_dir = tk.StringVar()
        self.dst_dir = tk.StringVar()
        self.max_width = tk.StringVar()  # blank means keep sizes

        # XLSX options
        self.make_xlsx = tk.BooleanVar(value=False)
        self.use_url_for_xlsx = tk.BooleanVar(value=False)
        self.url_base = tk.StringVar(value="")  # e.g., https://site.com/images/

        self._worker: Optional[threading.Thread] = None
        self._cancel = threading.Event()

        # progress text
        self.progress_text = tk.StringVar(value="Ready.")

        # destination auto-set tracking
        self._dst_was_auto_set = False

        self._build_ui()

        # react whenever source changes (browse, drag & drop)
        self.src_dir.trace_add("write", self._on_src_dir_change)

        self._update_xlsx_controls_state()

        # handle window close
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    # ---------------- UI construction ----------------
    def _build_ui(self):
        pad = 10

        frm = ttk.Frame(self, padding=pad)
        frm.pack(fill="both", expand=True)

        # one-line description
        desc = ttk.Label(
            frm,
            text="Recursively converts images to WEBP and copies videos into the destination folder.",
            wraplength=660,
        )
        desc.grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 10))

        row = 1

        # Source
        ttk.Label(frm, text="Source folder:").grid(row=row, column=0, sticky="w")
        ttk.Entry(frm, textvariable=self.src_dir, width=52).grid(
            row=row, column=1, sticky="we", padx=(5, 2)
        )
        src_btns = ttk.Frame(frm)
        src_btns.grid(row=row, column=2, sticky="e")
        ttk.Button(src_btns, text="Browse…", command=self.pick_src).pack(side="left")
        ttk.Button(src_btns, text="Open", command=self.open_src).pack(side="left", padx=(4, 0))

        # Source preview area (initially hidden)
        row += 1
        self.preview_frame = ttk.LabelFrame(frm, text="Source preview")
        self.preview_frame.grid(row=row, column=0, columnspan=3, sticky="we", pady=(4, 0))
        self.preview_count_label = ttk.Label(
            self.preview_frame, text="No valid source folder selected."
        )
        self.preview_count_label.pack(anchor="w", padx=6, pady=(4, 2))
        self.preview_list = tk.Listbox(self.preview_frame, height=4)
        self.preview_list.pack(fill="both", expand=True, padx=6, pady=(0, 4))

        # Drag & drop target on preview area (if available)
        if DND_AVAILABLE:
            self.preview_frame.drop_target_register(DND_FILES)
            self.preview_frame.dnd_bind("<<Drop>>", self._on_drop)

        # hide preview until we have a valid source
        self.preview_frame.grid_remove()

        # Destination
        row += 1
        ttk.Label(frm, text="Destination folder:").grid(row=row, column=0, sticky="w", pady=(8, 0))
        self.dst_entry = ttk.Entry(frm, textvariable=self.dst_dir, width=52)
        self.dst_entry.grid(row=row, column=1, sticky="we", padx=(5, 2), pady=(8, 0))

        # If user types into destination entry, consider it manually overridden.
        self.dst_entry.bind("<KeyRelease>", self._on_dst_manual_edit)

        dst_btns = ttk.Frame(frm)
        dst_btns.grid(row=row, column=2, sticky="e", pady=(8, 0))
        ttk.Button(dst_btns, text="Browse…", command=self.pick_dst).pack(side="left")
        ttk.Button(dst_btns, text="Open", command=self.open_dst).pack(side="left", padx=(4, 0))

        # Max width
        row += 1
        ttk.Label(frm, text="Max width (px):").grid(row=row, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(frm, textvariable=self.max_width, width=20).grid(
            row=row, column=1, sticky="w", padx=(5, 5), pady=(8, 0)
        )
        ttk.Label(frm, text="optional").grid(row=row, column=2, sticky="w", pady=(8, 0))

        # XLSX checkbox
        row += 1
        self.chk_make_xlsx = ttk.Checkbutton(
            frm,
            text="Create XLSX index of images (converted + copied WEBP)",
            variable=self.make_xlsx,
            command=self._update_xlsx_controls_state,
        )
        self.chk_make_xlsx.grid(row=row, column=0, columnspan=3, sticky="w", pady=(6, 0))

        # URL base + toggle
        row += 1
        url_row = ttk.Frame(frm)
        url_row.grid(row=row, column=0, columnspan=3, sticky="we", pady=(2, 0))
        self.chk_use_url = ttk.Checkbutton(
            url_row,
            text="Use URL base in XLSX (Column B becomes URL_BASE + filename.webp)",
            variable=self.use_url_for_xlsx,
            command=self._update_xlsx_controls_state,
        )
        self.chk_use_url.pack(side="left")

        # URL base entry
        row += 1
        ttk.Label(frm, text="URL base (optional):").grid(row=row, column=0, sticky="w", pady=(4, 0))
        self.entry_url_base = ttk.Entry(frm, textvariable=self.url_base, width=60)
        self.entry_url_base.grid(row=row, column=1, sticky="we", padx=(5, 5), pady=(4, 0))
        ttk.Label(frm, text="Example: https://mywebsite.com/product-images/").grid(
            row=row, column=2, sticky="w", pady=(4, 0)
        )

        # Buttons
        row += 1
        btns = ttk.Frame(frm)
        btns.grid(row=row, column=0, columnspan=3, sticky="we", pady=(12, 0))
        self.start_btn = ttk.Button(btns, text="Convert", command=self.start)
        self.start_btn.pack(side="left")
        self.cancel_btn = ttk.Button(btns, text="Cancel", command=self.cancel, state="disabled")
        self.cancel_btn.pack(side="left", padx=(8, 0))

        # Progress text label
        row += 1
        self.progress_label = ttk.Label(frm, textvariable=self.progress_text)
        self.progress_label.grid(row=row, column=0, columnspan=3, sticky="w", pady=(12, 0))

        # Progress bar
        row += 1
        self.prog = ttk.Progressbar(frm, orient="horizontal", mode="determinate")
        self.prog.grid(row=row, column=0, columnspan=3, sticky="we", pady=(4, 0))

        # Log
        row += 1
        self.log = tk.Text(frm, height=10, wrap="word")
        self.log.grid(row=row, column=0, columnspan=3, sticky="nsew", pady=(10, 0))
        frm.grid_columnconfigure(1, weight=1)

    # ---------------- Folder pickers / openers ----------------
    def pick_src(self):
        desktop = str(Path.home() / "Desktop")
        d = filedialog.askdirectory(title="Select source directory", initialdir=desktop)
        if d:
            self.src_dir.set(d)

    def pick_dst(self):
        desktop = str(Path.home() / "Desktop")
        d = filedialog.askdirectory(title="Select destination directory", initialdir=desktop)
        if d:
            self.dst_dir.set(d)
            self._dst_was_auto_set = False  # user chose manually

    def _on_dst_manual_edit(self, event=None):
        # User typed in the destination box => manual override
        self._dst_was_auto_set = False

    def open_src(self):
        path = self.src_dir.get().strip()
        if not path:
            messagebox.showinfo("Open source", "Source folder is empty.")
            return
        self._open_path_in_file_manager(path)

    def open_dst(self):
        path = self.dst_dir.get().strip()
        if not path:
            messagebox.showinfo("Open destination", "Destination folder is empty.")
            return
        self._open_path_in_file_manager(path)

    def _open_path_in_file_manager(self, path: str):
        p = Path(path).expanduser()
        if not p.exists():
            messagebox.showerror("Open folder", f"Path does not exist:\n{p}")
            return
        try:
            if sys.platform.startswith("darwin"):
                subprocess.run(["open", str(p)], check=False)
            elif os.name == "nt":
                os.startfile(str(p))  # type: ignore[attr-defined]
            else:
                subprocess.run(["xdg-open", str(p)], check=False)
        except Exception as e:
            messagebox.showerror("Open folder", f"Could not open folder:\n{e}")

    # ---------------- Source preview + drag & drop ----------------
    def _on_src_dir_change(self, *args):
        path_str = self.src_dir.get().strip()
        if not path_str:
            self._set_preview(None)
            return
        p = Path(path_str).expanduser()
        if not p.is_dir():
            self._set_preview(None)
            return

        # Update preview
        self._set_preview(p)

        # Auto-set destination AFTER source selected:
        # If destination is empty OR it was previously auto-set (keep auto behavior)
        if self.dst_dir.get().strip() == "" or self._dst_was_auto_set:
            auto_dst = make_auto_destination_for_source(p)
            self.dst_dir.set(str(auto_dst))
            self._dst_was_auto_set = True

    def _set_preview(self, path: Optional[Path]):
        if path is None:
            def _do_clear():
                self.preview_list.delete(0, "end")
                self.preview_count_label.config(text="No valid source folder selected.")
                self.preview_frame.grid_remove()
            self._call_on_main(_do_clear)
            return

        def _do_update():
            self.preview_list.delete(0, "end")
            try:
                images = [pp for pp in iter_files(path) if is_image(pp)]
            except Exception:
                images = []
            count = len(images)
            self.preview_count_label.config(
                text=f"Found {count} image file(s) in this folder (and subfolders)."
            )
            limit = 50
            for img_path in images[:limit]:
                try:
                    rel = img_path.relative_to(path)
                except ValueError:
                    rel = img_path
                self.preview_list.insert("end", str(rel))
            if count > limit:
                self.preview_list.insert(
                    "end", f"... and {count - limit} more image(s) not shown."
                )
            self.preview_frame.grid()
        self._call_on_main(_do_update)

    def _on_drop(self, event):
        """
        Handle files/folders dropped from Finder/Explorer onto the preview area.
        If it's a file, we take its parent directory as the source.
        If it's a folder, we use the folder directly.
        """
        data = event.data
        if not data:
            return
        try:
            paths = self.tk.splitlist(data)
        except Exception:
            paths = [data]
        if not paths:
            return
        first = paths[0]
        p = Path(first).expanduser()
        if p.is_file():
            p = p.parent
        if not p.is_dir():
            return
        self.src_dir.set(str(p))

    # ---------------- Start / cancel ----------------
    def start(self):
        if self._worker and self._worker.is_alive():
            return

        src = Path(self.src_dir.get()).expanduser()
        if not src.exists() or not src.is_dir():
            messagebox.showerror("Error", "Please select a valid source directory.")
            return

        dst_raw = self.dst_dir.get().strip()
        if not dst_raw:
            # Should not happen because we auto-set after source, but handle anyway.
            dst = make_auto_destination_for_source(src)
            self.dst_dir.set(str(dst))
            self._dst_was_auto_set = True
        else:
            dst = Path(dst_raw).expanduser()

        # Create destination if missing
        try:
            if dst.exists() and not dst.is_dir():
                messagebox.showerror(
                    "Error", f"Destination path exists and is not a directory:\n{dst}"
                )
                return
            dst.mkdir(parents=True, exist_ok=True)
        except Exception as e:
            messagebox.showerror("Error", f"Could not create destination directory:\n{e}")
            return

        if src.resolve() == dst.resolve():
            messagebox.showerror("Error", "Source and destination must be different.")
            return

        mw_raw = self.max_width.get().strip()
        if mw_raw == "":
            max_w = None  # keep sizes
        else:
            try:
                max_w_val = int(mw_raw)
                if max_w_val < 1:
                    raise ValueError
                max_w = max_w_val
            except ValueError:
                messagebox.showerror(
                    "Error", "Max width must be a positive integer or left blank."
                )
                return

        # XLSX + URL base validation
        want_xlsx = bool(self.make_xlsx.get())
        want_url = bool(self.use_url_for_xlsx.get())
        url_base = self.url_base.get().strip()

        if want_xlsx and want_url and url_base == "":
            messagebox.showerror(
                "Error", "URL base is empty, but 'Use URL base in XLSX' is checked."
            )
            return

        self.log_delete()
        self._cancel.clear()
        self.start_btn.config(state="disabled")
        self.cancel_btn.config(state="normal")
        self.prog.config(value=0, maximum=100)
        self.progress_text.set("Starting…")

        # kick off worker
        self._worker = threading.Thread(
            target=self._run_conversion,
            args=(src, dst, max_w, want_xlsx, want_url, url_base),
            daemon=True,
        )
        self._worker.start()
        self.after(150, self._poll_worker)

    def cancel(self):
        if self._worker and self._worker.is_alive():
            self._cancel.set()
            self.append_log("Cancelling…")

    def _poll_worker(self):
        if self._worker and self._worker.is_alive():
            self.after(150, self._poll_worker)
        else:
            self._set_buttons_idle()

    # ---------------- Conversion worker ----------------
    def _run_conversion(
        self,
        src: Path,
        dst: Path,
        max_width: Optional[int],
        want_xlsx: bool,
        want_url: bool,
        url_base: str,
    ):
        xlsx_rows: List[Tuple[str, str]] = []  # (filename_no_ext, ColumnB_value)
        xlsx_path: Optional[Path] = None

        # normalize URL base once
        if want_url and url_base:
            if not url_base.endswith("/"):
                url_base = url_base + "/"

        start_time = time.time()

        try:
            # Collect both images and videos for a single progress bar
            files = [p for p in iter_files(src) if is_image(p) or is_video(p)]
            total = len(files)
            converted = 0
            copied_videos = 0
            copied_webps = 0
            errors = 0

            if total == 0:
                self.append_log("No images or videos found.")
                self.set_progress(0, 1, "No images or videos found.")
                return

            self.set_progress(0, total, f"Found {total} files to process…")

            for src_path in files:
                if self._cancel.is_set():
                    self.append_log("Operation cancelled.")
                    break

                rel = src_path.relative_to(src)
                dst_path = dst / rel

                try:
                    if is_video(src_path):
                        # copy videos as-is
                        copy_file(src_path, dst_path)
                        copied_videos += 1
                        self.append_log(f"⧉ VIDEO: {src_path}  →  {dst_path}")

                    elif is_image(src_path):
                        if is_webp(src_path):
                            # copy existing webp without recompressing
                            copy_file(src_path, dst_path)
                            copied_webps += 1
                            self.append_log(f"⧉ WEBP COPY: {src_path}  →  {dst_path}")

                            # index copied webp too
                            if want_xlsx:
                                filename_no_ext = dst_path.stem  # Column A
                                if want_url and url_base:
                                    col_b = url_base + dst_path.name
                                else:
                                    col_b = str(dst_path.resolve())
                                xlsx_rows.append((filename_no_ext, col_b))
                        else:
                            # convert other image formats to .webp
                            dst_webp = dst_path.with_suffix(".webp")
                            convert_to_webp(src_path, dst_webp, max_width)
                            converted += 1
                            self.append_log(f"✓ IMAGE: {src_path}  →  {dst_webp}")

                            # record for XLSX
                            if want_xlsx:
                                filename_no_ext = dst_webp.stem  # Column A
                                if want_url and url_base:
                                    col_b = url_base + dst_webp.name
                                else:
                                    col_b = str(dst_webp.resolve())
                                xlsx_rows.append((filename_no_ext, col_b))

                except Exception as e:
                    errors += 1
                    self.append_log(f"✗ FAILED: {src_path}  —  {e}")

                processed = converted + copied_videos + copied_webps + errors
                status = (
                    f"Processed {processed} / {total} "
                    f"(images: {converted}, webp copied: {copied_webps}, "
                    f"videos: {copied_videos}, errors: {errors})"
                )
                self.set_progress(processed, total, status)

            # After processing: write xlsx if requested
            if want_xlsx:
                if xlsx_rows:
                    xlsx_path = dst / "converted_images.xlsx"
                    try:
                        write_xlsx(xlsx_rows, xlsx_path)
                        self.append_log(f"\nXLSX index saved: {xlsx_path}")
                        self.append_log(f"Rows written     : {len(xlsx_rows)}")
                    except Exception as e:
                        self.append_log(f"✗ Failed to write XLSX: {e}")
                else:
                    self.append_log("\nNo images to index; XLSX not created.")

            duration = time.time() - start_time

            self.append_log("\nDone.")
            self.append_log(f"Images converted : {converted}")
            self.append_log(f"WEBP copied      : {copied_webps}")
            self.append_log(f"Videos copied    : {copied_videos}")
            self.append_log(f"Errors           : {errors}")
            if xlsx_path:
                self.append_log(f"XLSX index       : {xlsx_path}")
            self.set_progress(total, max(total, 1), "Done.")

            self._show_summary(converted, copied_webps, copied_videos, errors, xlsx_path, duration)

        finally:
            self._set_buttons_idle()

    # --- UI helpers for thread-safe updates
    def append_log(self, text: str):
        def _do():
            self.log.insert("end", text + "\n")
            self.log.see("end")

        self._call_on_main(_do)

    def log_delete(self):
        def _do():
            self.log.delete("1.0", "end")

        self._call_on_main(_do)

    def set_progress(self, value: int, maximum: int, text: Optional[str] = None):
        def _do():
            self.prog.config(maximum=maximum, value=value)
            if text is not None:
                self.progress_text.set(text)

        self._call_on_main(_do)

    def _call_on_main(self, fn):
        if self._is_main_thread():
            fn()
        else:
            self.after(0, fn)

    @staticmethod
    def _is_main_thread():
        return threading.current_thread() == threading.main_thread()

    def _set_buttons_idle(self):
        def _do():
            self.start_btn.config(state="normal")
            self.cancel_btn.config(state="disabled")

        self._call_on_main(_do)

    # ---------------- XLSX controls state ----------------
    def _update_xlsx_controls_state(self):
        make = bool(self.make_xlsx.get())
        use_url = bool(self.use_url_for_xlsx.get())

        if not make:
            self.chk_use_url.config(state="disabled")
            self.use_url_for_xlsx.set(False)
            self.entry_url_base.config(state="disabled")
        else:
            self.chk_use_url.config(state="normal")
            self.entry_url_base.config(state="normal" if use_url else "disabled")

    # ---------------- Summary dialog ----------------
    def _show_summary(
        self,
        converted: int,
        copied_webps: int,
        copied_videos: int,
        errors: int,
        xlsx_path: Optional[Path],
        duration: float,
    ):
        lines = [
            f"Images converted : {converted}",
            f"WEBP copied      : {copied_webps}",
            f"Videos copied    : {copied_videos}",
            f"Errors           : {errors}",
            f"Time             : {duration:.1f} seconds",
        ]
        if xlsx_path:
            lines.append(f"XLSX index       : {xlsx_path}")
        msg = "\n".join(lines)

        def _do():
            messagebox.showinfo("Conversion finished", msg)

        self._call_on_main(_do)

    def on_close(self):
        if self._worker and self._worker.is_alive():
            if messagebox.askyesno(
                "Quit", "Conversion is still running. Do you want to quit anyway?"
            ):
                self._cancel.set()
                self.destroy()
        else:
            self.destroy()


if __name__ == "__main__":
    App().mainloop()
